import io
import logging
import sys
import tempfile
import unittest
from contextlib import redirect_stderr, redirect_stdout
from pathlib import Path
from unittest import mock

from openpyxl import load_workbook

from tests.fakes import CONFIG, FakeClient
import main as launcher
from Commands import portfolio_cycle, reports_download, scores_export
from constants import OPTION_VENDORS


class QuietTest(unittest.TestCase):
    def setUp(self):
        logging.disable(logging.CRITICAL)
        self.addCleanup(logging.disable, logging.NOTSET)
        for patcher in (
            mock.patch("logging.basicConfig"),
            redirect_stdout(io.StringIO()),
            redirect_stderr(io.StringIO()),
        ):
            patcher.__enter__()
            self.addCleanup(patcher.__exit__, None, None, None)


class PortfolioCycleArgsTest(QuietTest):
    def test_every_configured_cycle_is_accepted(self):
        for option in sorted(OPTION_VENDORS):
            with self.subTest(option=option), mock.patch.object(
                sys, "argv", ["portfolio_cycle.py", "--option", str(option), "--dry-run"]
            ):
                try:
                    args = portfolio_cycle.parse_args()
                except SystemExit:
                    self.fail(f"--option {option} is rejected by argparse")
                self.assertEqual(args.option, option)


class MenuTest(QuietTest):
    """The `python main.py --cli` menu must be able to start every command."""

    def run_menu_choice(self, choice: str, module) -> mock.Mock:
        load_config = mock.Mock(side_effect=RuntimeError("stop after argument parsing"))
        # menu choice, cycle, "press Enter", exit
        with mock.patch.object(sys, "argv", ["main.py"]), mock.patch(
            "builtins.input", side_effect=[choice, "3", "", "0"]
        ), mock.patch.object(module, "load_config", load_config):
            launcher.run_cli()
        return load_config

    def test_menu_download_reports_gets_past_argument_parsing(self):
        self.assertTrue(
            self.run_menu_choice("2", reports_download).called,
            "Download Reports exited in argparse",
        )

    def test_menu_update_score_history_gets_past_argument_parsing(self):
        self.assertTrue(
            self.run_menu_choice("3", scores_export).called,
            "Update Score History exited in argparse",
        )


class ScoresExportCommandTest(QuietTest):
    def run_export(self, cycle: int) -> list[tuple]:
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        history = Path(tmp.name) / "history.xlsx"
        client = FakeClient(
            [{"domain": "kalbe.co.id", "name": "Kalbe Farma"}, {"domain": "Zerpidio.com", "name": "Zerpidio"}],
            {"kalbe.co.id": (92, "A"), "zerpidio.com": (81, "B")},
        )
        argv = ["scores_export.py", "--cycle", str(cycle), "--history", str(history)]
        with mock.patch.object(sys, "argv", argv), mock.patch.object(
            scores_export, "load_config", return_value=CONFIG
        ), mock.patch.object(scores_export, "SecurityScorecardClient", return_value=client):
            scores_export.main()
        worksheet = load_workbook(history).active
        return [(row[1], row[3], row[4], row[5]) for row in worksheet.iter_rows(min_row=3, values_only=True)]

    def test_cycle_one_records_pinned_domains(self):
        self.assertEqual(
            self.run_export(1),
            [("kalbe.co.id", 92, "A", 1), ("Zerpidio.com", 81, "B", 1)],
        )

    def test_later_cycles_skip_pinned_domains(self):
        self.assertEqual(self.run_export(2), [("Zerpidio.com", 81, "B", 2)])


if __name__ == "__main__":
    unittest.main()
