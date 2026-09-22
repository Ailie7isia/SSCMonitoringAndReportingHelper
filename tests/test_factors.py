import logging
import tempfile
import unittest
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

from tests.fakes import FakeClient
from models import Company
from Services.factors import (
    FACTORS,
    HEADERS,
    HISTORY_SOURCE,
    LIVE_SOURCE,
    FactorSnapshot,
    append_factor_history,
    grade_for_score,
    latest_by_month,
    read_factor_history,
    snapshot_from_factors,
    snapshots_from_history,
    update_factor_history,
)
from Services.scores import FACTOR_SHEET, append_score_history, score_trends_from_history


def utc_now() -> datetime:
    return datetime.now(timezone.utc).replace(tzinfo=None)


def months_ago(count: int) -> tuple[int, int]:
    now = utc_now()
    year, month = now.year, now.month - count
    while month <= 0:
        year, month = year - 1, month + 12
    return year, month


def iso_months_ago(count: int, day: int = 15) -> str:
    year, month = months_ago(count)
    return f"{year:04d}-{month:02d}-{day:02d}T00:00:00.000Z"


def snapshot(domain="kalbe.co.id", when=None, score=95.0, source=LIVE_SOURCE, issues=None, cycle=1):
    return FactorSnapshot(
        retrieved=when or utc_now(),
        domain=domain,
        company="Kalbe Farma",
        source=source,
        cycle=cycle,
        scores={key: score for key in FACTORS},
        issues={key: issues for key in FACTORS} if issues is not None else {},
    )


ALL_FACTORS = {key: (95, 0) for key in FACTORS}


class FactorTestBase(unittest.TestCase):
    def setUp(self):
        logging.disable(logging.CRITICAL)
        self.addCleanup(logging.disable, logging.NOTSET)
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        self.path = Path(tmp.name) / "history.xlsx"


class ParsingTest(FactorTestBase):
    def test_live_snapshot_reads_scores_and_counts_only_real_issue_types(self):
        entries = [
            {
                "name": "application_security",
                "score": 69,
                "issue_summary": [
                    {"severity": "medium", "type": "csp_no_policy", "count": 4},
                    {"severity": "low", "type": "hsts_incorrect", "count": 1},
                    {"severity": "positive", "type": "tls_ok", "count": 9},
                    {"severity": "info", "type": "note", "count": 1},
                ],
            },
            {"name": "dns_health", "score": 98, "issue_summary": []},
            {"name": "brand_new_factor", "score": 10},
        ]
        result = snapshot_from_factors(Company("Hydrococo.ID", "Hydrococo"), entries, cycle=4, retrieved=utc_now())
        self.assertEqual(result.domain, "hydrococo.id")
        self.assertEqual(result.scores, {"application_security": 69.0, "dns_health": 98.0})
        self.assertEqual(result.issues, {"application_security": 2, "dns_health": 0})
        self.assertEqual((result.source, result.cycle), (LIVE_SOURCE, 4))

    def test_response_without_scores_gives_no_snapshot(self):
        self.assertIsNone(snapshot_from_factors(Company("a.com", "A"), [], cycle=1, retrieved=utc_now()))

    def test_history_entries_become_rounded_monthly_snapshots(self):
        entries = [{
            "date": "2026-08-14T00:00:00.000Z",
            "factors": [
                {"name": "application_security", "score": 99.99999999999997},
                {"name": "dns_health", "score": 98.36666666666663},
            ],
        }, {"date": "not a date", "factors": []}]
        [result] = snapshots_from_history(Company("kalbe.co.id", "Kalbe Farma"), entries)
        self.assertEqual(result.retrieved, datetime(2026, 8, 14))
        self.assertEqual(result.scores, {"application_security": 100.0, "dns_health": 98.4})
        self.assertEqual(result.source, HISTORY_SOURCE)
        self.assertEqual(result.issues, {})

    def test_grades_follow_securityscorecard_bands(self):
        self.assertEqual(
            [grade_for_score(score) for score in (100, 90, 89.9, 80, 70, 60, 59, None)],
            ["A", "A", "B", "B", "C", "D", "F", "?"],
        )

    def test_latest_by_month_keeps_the_newest_snapshot(self):
        older = snapshot(when=datetime(2026, 9, 3), score=80)
        newer = snapshot(when=datetime(2026, 9, 20), score=90)
        august = snapshot(when=datetime(2026, 8, 30), score=70)
        result = latest_by_month([newer, older, august])
        self.assertEqual(result["kalbe.co.id"][(2026, 9)], newer)
        self.assertEqual(result["kalbe.co.id"][(2026, 8)], august)


class WorkbookTest(FactorTestBase):
    def test_append_and_read_round_trip(self):
        saved = [snapshot(issues=3), snapshot("zerpidio.com", when=datetime(2026, 7, 18), score=88.4, source=HISTORY_SOURCE, cycle=None)]
        self.assertEqual(append_factor_history(saved, self.path), 2)

        worksheet = load_workbook(self.path)[FACTOR_SHEET]
        self.assertEqual([cell.value for cell in worksheet[2]], HEADERS)
        self.assertIn("Information Leak", HEADERS)
        self.assertIn("Information Leak issues", HEADERS)

        loaded = sorted(read_factor_history(self.path), key=lambda item: item.domain)
        self.assertEqual([item.domain for item in loaded], ["kalbe.co.id", "zerpidio.com"])
        self.assertEqual(loaded[0].scores["network_security"], 95.0)
        self.assertEqual(loaded[0].issues["network_security"], 3)
        self.assertEqual(loaded[1].scores["network_security"], 88.4)
        self.assertIsNone(loaded[1].issues["network_security"])
        self.assertEqual(loaded[1].source, HISTORY_SOURCE)

    def test_nothing_to_append_leaves_no_file(self):
        self.assertEqual(append_factor_history([], self.path), 0)
        self.assertFalse(self.path.exists())

    def test_factor_sheet_does_not_disturb_score_history(self):
        append_score_history([Company("kalbe.co.id", "Kalbe Farma", "A", 96)], 1, self.path)
        append_factor_history([snapshot()], self.path)
        # Excel saves whichever tab was selected as the active sheet.
        workbook = load_workbook(self.path)
        workbook.active = workbook.sheetnames.index(FACTOR_SHEET)
        workbook.save(self.path)

        self.assertIn("kalbe.co.id", score_trends_from_history(self.path))
        append_score_history([Company("kalbe.co.id", "Kalbe Farma", "A", 97)], 1, self.path)
        workbook = load_workbook(self.path)
        score_rows = [row[3] for row in workbook.worksheets[0].iter_rows(min_row=3, values_only=True)]
        self.assertEqual(score_rows, [96, 97])
        self.assertEqual(workbook[FACTOR_SHEET].max_row, 3)

    def test_score_sheet_created_after_factor_sheet_comes_first(self):
        append_factor_history([snapshot()], self.path)
        append_score_history([Company("kalbe.co.id", "Kalbe Farma", "A", 96)], 1, self.path)
        workbook = load_workbook(self.path)
        self.assertNotEqual(workbook.sheetnames[0], FACTOR_SHEET)
        self.assertIn("kalbe.co.id", score_trends_from_history(self.path))
        self.assertEqual(len(read_factor_history(self.path)), 1)

    def test_unreadable_workbook_reads_as_no_history(self):
        self.path.write_bytes(b"not an xlsx")
        self.assertEqual(read_factor_history(self.path), [])


class UpdateTest(FactorTestBase):
    def client(self):
        client = FakeClient()
        client.factors = {"kalbe.co.id": ALL_FACTORS, "zerpidio.com": ALL_FACTORS}
        return client

    def test_backfills_missing_months_and_adds_a_live_snapshot(self):
        client = self.client()
        client.factor_history = {"kalbe.co.id": [
            (iso_months_ago(14), {key: 70.0 for key in FACTORS}),  # older than the backfill window
            (iso_months_ago(2), {key: 90.0 for key in FACTORS}),
            (iso_months_ago(1), {key: 92.0 for key in FACTORS}),
            (iso_months_ago(0, day=1), {key: 93.0 for key in FACTORS}),  # current month comes from live
        ]}
        result = update_factor_history(client, [Company("kalbe.co.id", "Kalbe Farma")], 1, self.path)
        self.assertEqual((result.live, result.backfilled, result.failed), (1, 2, []))

        saved = latest_by_month(read_factor_history(self.path))["kalbe.co.id"]
        self.assertEqual(set(saved), {months_ago(2), months_ago(1), months_ago(0)})
        self.assertEqual(saved[months_ago(0)].source, LIVE_SOURCE)
        self.assertEqual(saved[months_ago(0)].cycle, 1)
        self.assertEqual(saved[months_ago(1)].scores["dns_health"], 92.0)

    def test_does_not_call_history_again_once_past_months_are_saved(self):
        past = [
            snapshot(when=datetime(*months_ago(count), 15), source=HISTORY_SOURCE, cycle=None)
            for count in range(1, 13)
        ]
        append_factor_history(past, self.path)
        client = self.client()
        result = update_factor_history(client, [Company("kalbe.co.id", "Kalbe Farma")], 1, self.path)
        self.assertEqual(client.factor_history_calls, [])
        self.assertEqual((result.live, result.backfilled), (1, 0))

    def test_backfill_can_be_turned_off(self):
        client = self.client()
        update_factor_history(client, [Company("kalbe.co.id", "Kalbe Farma")], 1, self.path, backfill=False)
        self.assertEqual(client.factor_history_calls, [])

    def test_one_failing_domain_does_not_stop_the_others(self):
        client = self.client()
        client.fail_factors = {"kalbe.co.id"}
        result = update_factor_history(
            client, [Company("kalbe.co.id", "Kalbe Farma"), Company("zerpidio.com", "Zerpidio")], 2, self.path
        )
        self.assertEqual((result.live, result.failed), (1, ["kalbe.co.id"]))
        self.assertEqual([item.domain for item in read_factor_history(self.path)], ["zerpidio.com"])

    def test_domain_without_factor_scores_is_reported(self):
        result = update_factor_history(FakeClient(), [Company("new.example", "New")], 1, self.path)
        self.assertEqual((result.live, result.failed), (0, ["new.example"]))
        self.assertFalse(self.path.exists())


if __name__ == "__main__":
    unittest.main()
