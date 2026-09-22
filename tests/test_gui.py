"""Headless-ish smoke tests for the customtkinter dashboard (window stays withdrawn)."""

import logging
import queue
import tempfile
import time
import tkinter as tk
import unittest
from datetime import datetime, timedelta, timezone
from pathlib import Path
from unittest import mock

from openpyxl import load_workbook

from tests.fakes import FakeClient
from tests.test_scores import write_history
from models import Company
from Services.factors import FACTORS, HISTORY_SOURCE, LIVE_SOURCE, FactorSnapshot, append_factor_history, read_factor_history
from Services.portfolio import target_domains
from Services.scores import HistoricalScoreTrend
from Services.rate_limit import RateLimitStatus
from ssc_client import RateLimitError

try:
    import customtkinter as ctk
    import gui
except Exception as exc:  # pragma: no cover - environment without Tk
    gui = None
    IMPORT_ERROR = exc


def descendants(widget):
    for child in widget.winfo_children():
        yield child
        yield from descendants(child)


def label_texts(widget) -> list[str]:
    return [w.cget("text") for w in descendants(widget) if isinstance(w, ctk.CTkLabel)]


def all_texts(widget) -> list[str]:
    """Text of every label, including the plain Tk labels used for factor tables."""
    return [w.cget("text") for w in descendants(widget) if isinstance(w, tk.Label)]


def utc_now() -> datetime:
    return datetime.now(timezone.utc).replace(tzinfo=None)


def last_month() -> datetime:
    return (utc_now().replace(day=1) - timedelta(days=1)).replace(day=15, hour=0, minute=0, second=0, microsecond=0)


def factor_snapshot(domain="kalbe.co.id", when=None, score=95.0, source=LIVE_SOURCE, issues=None):
    return FactorSnapshot(
        retrieved=when or utc_now(),
        domain=domain,
        company=domain,
        source=source,
        cycle=1,
        scores={key: score for key in FACTORS},
        issues={key: issues for key in FACTORS} if issues is not None else {},
    )


def no_network():
    raise AssertionError("test attempted to contact SecurityScorecard")


class InlineThread:
    """Runs a worker synchronously: Tk's after() needs mainloop() when called from another thread."""

    def __init__(self, target, daemon=None, **_kwargs):
        self.target = target

    def start(self):
        self.target()


class DashboardTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        if gui is None:
            raise unittest.SkipTest(f"GUI unavailable: {IMPORT_ERROR}")
        with mock.patch.object(gui.Dashboard, "refresh_dashboard"):
            try:
                cls.app = gui.Dashboard()
            except Exception as exc:
                raise unittest.SkipTest(f"Tk unavailable: {exc}")
        cls.app.withdraw()
        # Guard rails: nothing in this class may reach the real API.
        cls.app.refresh_dashboard = lambda: None
        cls.app._client = no_network

    @classmethod
    def tearDownClass(cls):
        cls.app.destroy()

    def setUp(self):
        self.errors: list[BaseException] = []
        self.app.report_callback_exception = lambda exc, value, tb: self.errors.append(value)
        threads = mock.patch.object(gui.threading, "Thread", InlineThread)
        threads.start()
        self.addCleanup(threads.stop)
        if gui.SCORE_HISTORY_PATH.exists():
            gui.SCORE_HISTORY_PATH.unlink()
        # Messages logged by an earlier test must not land in this test's Activity panel.
        while not self.app.log_queue.empty():
            self.app.log_queue.get_nowait()
        self.app.active_cycle = None
        self.app.show_dashboard()

    def wait_until_idle(self, timeout: float = 15) -> None:
        deadline = time.monotonic() + timeout
        while time.monotonic() < deadline:
            self.app.update()
            if not self.app.busy:
                return
            time.sleep(0.02)
        self.fail("background task did not finish")

    def clear_activity(self):
        self.app.activity.configure(state="normal")
        self.app.activity.delete("1.0", "end")
        self.app.activity.configure(state="disabled")

    # Activity log -----------------------------------------------------------------

    def test_genuine_failure_is_highlighted(self):
        self.clear_activity()
        self.app.log_queue.put("✗ Operation failed: boom")
        self.app._poll_log_queue()
        self.assertNotEqual(self.app.activity._textbox.tag_ranges("failure"), ())

    def test_service_error_logs_are_marked_as_failures(self):
        destination = queue.Queue()
        handler = gui.QueueLogHandler(destination)
        handler.emit(logging.LogRecord("ssc", logging.ERROR, __file__, 1, "Failed adding %s", ("a.com",), None))
        handler.emit(logging.LogRecord("ssc", logging.INFO, __file__, 1, "Added %s", ("b.com",), None))
        self.assertEqual([destination.get_nowait(), destination.get_nowait()], ["✗ Failed adding a.com", "Added b.com"])

    def test_successful_cycle_summary_is_not_highlighted_as_failure(self):
        self.clear_activity()
        self.app.log_queue.put("✓ Portfolio cycle complete: 2 added, 1 removed, 0 failed.")
        self.app._poll_log_queue()
        self.assertEqual(
            self.app.activity._textbox.tag_ranges("failure"),
            (),
            "a success message is rendered red because it contains ' failed'",
        )

    # Rate limits -------------------------------------------------------------------

    def test_rate_limit_countdown_shows_time_left_and_clears(self):
        self.addCleanup(self.app._show_rate_limit_countdown, None)
        self.app._show_rate_limit_countdown(datetime.now().astimezone() + timedelta(seconds=90))
        self.assertEqual(self.app.rate_limit_status.winfo_manager(), "pack")
        self.assertRegex(self.app.rate_limit_status.cget("text"), r"Resumes in 1:(29|30) \(at \d\d:\d\d:\d\d\)")

        self.app._show_rate_limit_countdown(None)
        self.assertEqual(self.app.rate_limit_status.winfo_manager(), "")

    def test_rate_limited_operation_starts_countdown(self):
        self.addCleanup(self.app._show_rate_limit_countdown, None)

        def task():
            raise RateLimitError("Too many requests", retry_after=120)

        self.app._run("Testing rate limit", task)
        self.wait_until_idle()
        self.app.update()
        remaining = (self.app._rate_limit_resume_at - datetime.now().astimezone()).total_seconds()
        self.assertAlmostEqual(remaining, 120, delta=5)
        self.assertEqual(self.app.rate_limit_status.winfo_manager(), "pack")

    def status(self, state="running", seconds=0, **overrides):
        values = dict(
            state=state,
            endpoint="reports/detailed",
            resume_at=datetime.now().astimezone() + timedelta(seconds=seconds) if seconds else None,
            spacing=3.0,
            requests=12,
            rate_limited=1,
            waited_on_limits=42.0,
            waited_on_pacing=9.0,
        )
        values.update(overrides)
        return RateLimitStatus(**values)

    def test_download_timer_counts_down_the_gap_and_shows_request_counts(self):
        self.addCleanup(self.app._show_rate_limit_status, None)
        self.app._show_rate_limit_status(self.status("pacing", seconds=5))
        text = self.app.rate_limit_status.cget("text")
        self.assertRegex(text, r"Next request in 0:0[45]")
        self.assertIn("12 sent · 1 limited · gap 3s", text)
        self.assertEqual(self.app.rate_limit_status.winfo_manager(), "pack")

    def test_download_timer_counts_down_a_limit(self):
        self.addCleanup(self.app._show_rate_limit_status, None)
        self.app._show_rate_limit_status(self.status("cooling", seconds=90, estimated=True))
        text = self.app.rate_limit_status.cget("text")
        self.assertRegex(text, r"Resumes in 1:(29|30) \(at \d\d:\d\d:\d\d\)")
        self.assertIn("Estimated", text)

    def test_download_timer_stays_between_waits_and_hides_when_finished(self):
        self.app._show_rate_limit_status(self.status("running"))
        self.assertIn("12 sent · 1 limited", self.app.rate_limit_status.cget("text"))
        self.assertEqual(self.app.rate_limit_status.winfo_manager(), "pack")
        self.app._show_rate_limit_status(self.status("finished"))
        self.assertEqual(self.app.rate_limit_status.winfo_manager(), "")

    def test_download_passes_pacer_status_to_the_timer(self):
        client = FakeClient(
            portfolio=[{"domain": "kalbe.co.id", "name": "Kalbe Farma"}],
            scores={"kalbe.co.id": (92, "A")},
        )
        client.create_rate_limits = 1
        client.rate_limit_retry_after = 30
        shown = []
        with tempfile.TemporaryDirectory() as tmp, \
                mock.patch.object(self.app, "_client", return_value=(client, "p1")), \
                mock.patch.object(gui, "REPORTS_DIR", Path(tmp)), \
                mock.patch.object(gui.threading, "Thread", InlineThread), \
                mock.patch("Services.reports.time.sleep"), \
                mock.patch.object(self.app, "_show_rate_limit_status", side_effect=shown.append):
            self.app.active_cycle = None
            self.app._start_reports_download("detailed_report", "Detailed PDF reports")
            self.wait_until_idle()
        states = [status.state for status in shown if status is not None]
        self.assertIn("cooling", states)
        self.assertEqual(states[-1], "finished")

    # Pages -------------------------------------------------------------------------

    def test_cycle_label_matches_configured_cycle(self):
        companies = [Company(domain, domain) for domain in target_domains(3)]
        self.assertTrue(self.app._current_cycle_label(companies).startswith("Cycle 3 — "))
        self.assertEqual(self.app.active_cycle, 3)

    def test_score_trend_search_filters_rows(self):
        self.app.current_companies = [Company("kalbe.co.id", "Kalbe", "A", 91), Company("zerpidio.com", "Zerpidio", "C", 72)]
        self.app._show_score_trends({
            "kalbe.co.id": HistoricalScoreTrend(91.0, 85.0, 6.0, None),
            "zerpidio.com": HistoricalScoreTrend(72.0, 70.0, 2.0, None),
        })
        page = self.app.trends_page
        self.assertIn("zerpidio.com", label_texts(page))

        search = next(w for w in descendants(page) if isinstance(w, ctk.CTkEntry))
        search.insert(0, "kalbe")
        self.app.update_idletasks()

        self.assertEqual(self.errors, [], "typing in the search box raised inside a Tk callback")
        texts = label_texts(page)
        self.assertIn("kalbe.co.id", texts)
        self.assertNotIn("zerpidio.com", texts)

    def test_report_page_renders_workbook_months(self):
        this_month = datetime.now(timezone.utc).replace(tzinfo=None, day=1, hour=0, minute=0, second=0, microsecond=0)
        last_month = this_month - timedelta(days=3)
        write_history(gui.SCORE_HISTORY_PATH, [
            (last_month, "kalbe.co.id", "Kalbe", 85, "B", 1),
            (this_month, "kalbe.co.id", "Kalbe", 91, "A", 1),
            (this_month, "zerpidio.com", "Zerpidio", 65, None, 2),
        ])
        snapshots = self.app._history_by_month()
        self.assertEqual(list(snapshots), [this_month.strftime("%B %Y"), last_month.strftime("%B %Y")])
        self.app._show_report_page(snapshots)
        self.app.update_idletasks()
        self.assertEqual(self.errors, [])
        texts = label_texts(self.app.report_page)
        self.assertIn("kalbe.co.id", texts)
        self.assertIn("zerpidio.com", texts)

    def test_latest_issue_report_prefers_newest_dated_folder(self):
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        root = Path(tmp.name)
        for day, grade in (("2026-09-03", "B"), ("2026-09-10", "A")):
            folder = root / day / "issues"
            folder.mkdir(parents=True)
            (folder / f"[{grade}] - Kalbe Farma - Issue Report - September 2026.csv").write_text("x")
        with mock.patch.object(gui, "REPORTS_DIR", root):
            found = self.app._latest_issue_report(Company("kalbe.co.id", "Kalbe Farma"))
        self.assertEqual(found.parent.parent.name, "2026-09-10")

    def test_latest_issue_report_finds_domain_suffixed_filename(self):
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        folder = Path(tmp.name) / "2026-09-14" / "issues"
        folder.mkdir(parents=True)
        for domain in ("kalventis.co.id", "kalventis.com"):
            (folder / f"[A] - Kalventis ({domain}) - Issue Report - September 2026.csv").write_text("x")
        with mock.patch.object(gui, "REPORTS_DIR", Path(tmp.name)):
            found = self.app._latest_issue_report(Company("kalventis.com", "Kalventis"))
        self.assertEqual(found.name, "[A] - Kalventis (kalventis.com) - Issue Report - September 2026.csv")

    # Factors -------------------------------------------------------------------------

    def test_factor_page_shows_every_domain_for_the_selected_month(self):
        self.app.current_companies = []
        self.app._show_factor_page([
            factor_snapshot("kalbe.co.id", score=95),
            factor_snapshot("zerpidio.com", score=72),
        ])
        self.app.update_idletasks()
        self.assertEqual(self.errors, [])
        texts = all_texts(self.app.factors_page)
        self.assertIn("kalbe.co.id", texts)
        self.assertIn("zerpidio.com", texts)
        self.assertIn("95", texts)
        self.assertIn("APP SEC", texts)
        self.assertIn("INFO LEAK", texts)
        links = [w for w in descendants(self.app.factors_page) if isinstance(w, tk.Label) and w.cget("cursor") == "hand2"]
        self.assertEqual(sorted(link.cget("text") for link in links), ["kalbe.co.id", "zerpidio.com"])

    def test_factor_page_active_filter_uses_the_refreshed_portfolio(self):
        self.app.current_companies = [Company("kalbe.co.id", "Kalbe Farma", "A", 96)]
        self.addCleanup(setattr, self.app, "current_companies", [])
        self.app._show_factor_page([factor_snapshot("kalbe.co.id"), factor_snapshot("zerpidio.com")])
        texts = all_texts(self.app.factors_page)
        self.assertIn("kalbe.co.id", texts)
        self.assertNotIn("zerpidio.com", texts)

    def test_factor_page_without_history_explains_how_to_fill_it(self):
        self.app._show_factor_page([])
        self.assertTrue(any("No factor history yet" in text for text in label_texts(self.app.factors_page)))

    def test_kalbe_view_shows_cards_changes_issues_and_monthly_history(self):
        self.app._show_domain_factors("kalbe.co.id", [
            factor_snapshot(when=last_month(), score=90, source=HISTORY_SOURCE),
            factor_snapshot(score=95, issues=3),
            factor_snapshot("zerpidio.com", score=60),
        ])
        self.app.update_idletasks()
        self.assertEqual(self.errors, [])
        texts = all_texts(self.app.factors_page)
        self.assertIn("kalbe.co.id", texts)
        self.assertIn("APPLICATION SECURITY", texts)
        self.assertIn("INFORMATION LEAK", texts)
        self.assertIn("3 issues", texts)
        self.assertTrue(any(text.startswith("↑ +5 vs ") for text in texts))
        self.assertIn(f"{last_month():%b %y}*", texts, "backfilled months are marked")
        self.assertNotIn("60", texts, "another domain's scores leaked into the view")

    def test_domain_view_without_history_explains_how_to_fill_it(self):
        self.app._show_domain_factors("kalbe.co.id", [])
        self.assertTrue(any("No factor history for kalbe.co.id" in text for text in label_texts(self.app.factors_page)))

    def test_kalbe_button_loads_factor_history_from_the_workbook(self):
        append_factor_history([factor_snapshot(score=88)], gui.SCORE_HISTORY_PATH)
        self.app.open_kalbe_factors()
        self.wait_until_idle()
        self.app.update()
        self.assertIsNotNone(self.app.factors_page)
        self.assertIn("88", all_texts(self.app.factors_page))

    def test_leaving_a_factor_page_restores_the_dashboard(self):
        self.app._show_factor_page([factor_snapshot()])
        self.app.show_dashboard()
        self.assertIsNone(self.app.factors_page)
        self.assertEqual(self.app.dashboard_header.winfo_manager(), "grid")

    # Background workflows -------------------------------------------------------------

    def test_declined_cycle_makes_no_changes(self):
        client = FakeClient([{"domain": "kalbe.co.id"}, {"domain": "old.example"}])
        with mock.patch.object(self.app, "_client", return_value=(client, "p")), mock.patch.object(
            self.app, "_ask_confirmation", return_value=False
        ):
            self.app._run_cycle(3)
            self.wait_until_idle()
        self.assertEqual((client.added, client.removed), ([], []))

    def test_confirmed_cycle_applies_plan(self):
        client = FakeClient([{"domain": d} for d in target_domains(3)] + [{"domain": "old.example"}])
        with mock.patch.object(self.app, "_client", return_value=(client, "p")), mock.patch.object(
            self.app, "_ask_confirmation", return_value=True
        ), mock.patch("Services.portfolio.time.sleep"):
            self.app._run_cycle(3)
            self.wait_until_idle()
        self.assertEqual((client.added, client.removed), ([], ["old.example"]))

    def test_score_export_for_later_cycle_skips_pinned_domains(self):
        client = FakeClient(
            [{"domain": "kalbe.co.id", "name": "Kalbe Farma"}, {"domain": "zerpidio.com", "name": "Zerpidio"}],
            {"kalbe.co.id": (92, "A"), "zerpidio.com": (81, "B")},
        )
        client.factors = {"zerpidio.com": {key: (81, 0) for key in FACTORS}}
        self.app.active_cycle = 2
        with mock.patch.object(self.app, "_client", return_value=(client, "p")):
            self.app._export_score_history()
            self.wait_until_idle()
        worksheet = load_workbook(gui.SCORE_HISTORY_PATH).active
        rows = [(r[1], r[3], r[4], r[5]) for r in worksheet.iter_rows(min_row=3, values_only=True)]
        self.assertEqual(rows, [("zerpidio.com", 81, "B", 2)])

    def test_score_export_also_saves_factor_scores(self):
        client = FakeClient(
            [{"domain": "kalbe.co.id", "name": "Kalbe Farma"}],
            {"kalbe.co.id": (92, "A")},
        )
        client.factors = {"kalbe.co.id": {key: (97, 1) for key in FACTORS}}
        client.factor_history = {"kalbe.co.id": [(last_month().strftime("%Y-%m-%dT00:00:00.000Z"), {key: 91.0 for key in FACTORS})]}
        self.app.active_cycle = 1
        with mock.patch.object(self.app, "_client", return_value=(client, "p")):
            self.app._export_score_history()
            self.wait_until_idle()
        saved = sorted(read_factor_history(gui.SCORE_HISTORY_PATH), key=lambda item: item.retrieved)
        self.assertEqual([(item.source, item.scores["dns_health"]) for item in saved], [(HISTORY_SOURCE, 91.0), (LIVE_SOURCE, 97.0)])
        self.assertEqual(saved[-1].issues["dns_health"], 1)


if __name__ == "__main__":
    unittest.main()
