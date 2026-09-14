import tempfile
import unittest
from datetime import datetime, timedelta, timezone
from pathlib import Path

from openpyxl import Workbook, load_workbook

import tests  # noqa: F401  (isolates the history workbook path)
from models import Company
from Services.scores import (
    HISTORY_HEADERS,
    HistoricalScoreTrend,
    append_score_history,
    current_month_history_status,
    score_change_over_past_month,
    score_trends_from_history,
)


def utc_now() -> datetime:
    return datetime.now(timezone.utc).replace(tzinfo=None)


def month_starts() -> tuple[datetime, datetime, datetime]:
    this_month = utc_now().replace(day=1, hour=0, minute=0, second=0, microsecond=0)
    last_month = (this_month - timedelta(days=1)).replace(day=1)
    last_year = this_month.replace(year=this_month.year - 1)
    return this_month, last_month, last_year


def write_history(path: Path, rows: list[tuple]) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet["A1"] = "SSC Monitoring Helper — Score History"
    worksheet.append(HISTORY_HEADERS)
    for row in rows:
        worksheet.append(list(row))
    workbook.save(path)


class ScoreHistoryTestBase(unittest.TestCase):
    def setUp(self):
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        self.path = Path(tmp.name) / "history.xlsx"


class AppendScoreHistoryTest(ScoreHistoryTestBase):
    def test_creates_workbook_with_header_and_sorted_rows(self):
        companies = [
            Company("zerpidio.com", "Zerpidio", "b", 81),
            Company("Bifarma.co.id", "Bifarma", "", None),
        ]
        self.assertEqual(append_score_history(companies, 2, self.path), 2)

        worksheet = load_workbook(self.path).active
        self.assertEqual(worksheet["A1"].value, "SSC Monitoring Helper — Score History")
        self.assertIn("A1:F1", {str(r) for r in worksheet.merged_cells.ranges})
        self.assertEqual([c.value for c in worksheet[2]], HISTORY_HEADERS)
        rows = [row[1:] for row in worksheet.iter_rows(min_row=3, values_only=True)]
        self.assertEqual(
            rows,
            [("Bifarma.co.id", "Bifarma", None, "UNKNOWN", 2), ("zerpidio.com", "Zerpidio", 81, "B", 2)],
        )

    def test_second_append_keeps_single_header(self):
        append_score_history([Company("a.com", "A", "A", 95)], 1, self.path)
        append_score_history([Company("a.com", "A", "A", 96)], 1, self.path)
        worksheet = load_workbook(self.path).active
        self.assertEqual(worksheet.max_row, 4)
        self.assertEqual([c.value for c in worksheet[2]], HISTORY_HEADERS)

    def test_upgrades_legacy_five_column_workbook(self):
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.merge_cells("A1:E1")
        worksheet["A1"] = "Legacy"
        for column, header in enumerate(HISTORY_HEADERS[:5], start=1):
            worksheet.cell(row=2, column=column, value=header)
        worksheet.append([utc_now(), "a.com", "A", 90, "A"])
        workbook.save(self.path)

        append_score_history([Company("a.com", "A", "A", 91)], 3, self.path)
        worksheet = load_workbook(self.path).active
        self.assertEqual(worksheet["F2"].value, "Cycle")
        self.assertIn("A1:F1", {str(r) for r in worksheet.merged_cells.ranges})
        self.assertEqual(worksheet.cell(row=4, column=6).value, 3)

    def test_empty_export_does_not_create_workbook(self):
        with self.assertRaises(ValueError):
            append_score_history([], 1, self.path)
        self.assertFalse(self.path.exists())


class TrendTest(ScoreHistoryTestBase):
    def test_month_and_year_trends(self):
        this_month, last_month, last_year = month_starts()
        write_history(self.path, [
            (last_month + timedelta(days=2), "Kalbe.co.id", "Kalbe", 80, "B", 1),
            (last_month + timedelta(days=5), "kalbe.co.id", "Kalbe", 85, "B", 1),
            (this_month, "kalbe.co.id", "Kalbe", 91, "A", 1),
            (last_year + timedelta(days=1), "kalbe.co.id", "Kalbe", 70, "C", 1),
            (last_month + timedelta(days=1), "zerpidio.com", "Zerpidio", 75, "C", 2),
            (this_month, "zerpidio.com", "Zerpidio", "n/a", "C", 2),
        ])
        trends = score_trends_from_history(self.path)
        self.assertEqual(trends["kalbe.co.id"], HistoricalScoreTrend(91.0, 85.0, 6.0, 21.0))
        self.assertEqual(trends["zerpidio.com"], HistoricalScoreTrend(None, 75.0, None, None))

    def test_current_month_status_counts_and_filters_by_cycle(self):
        this_month, last_month, _ = month_starts()
        newest = this_month + timedelta(minutes=5)
        write_history(self.path, [
            (this_month, "a.com", "A", 90, "A", 2),
            (newest, "b.com", "B", 80, "B", 2),
            (this_month, "c.com", "C", 70, "C", 3),
            (last_month, "d.com", "D", 60, "D", 2),
        ])
        self.assertEqual(current_month_history_status(self.path), (newest, 3))
        self.assertEqual(current_month_history_status(self.path, cycle=2), (newest, 2))

    def test_score_change_over_past_month_uses_newest_baseline_older_than_30_days(self):
        now = utc_now()
        write_history(self.path, [
            (now - timedelta(days=40), "kalbe.co.id", "Kalbe", 80, "B", 1),
            (now - timedelta(days=35), "kalbe.co.id", "Kalbe", 82, "B", 1),
            (now - timedelta(days=10), "kalbe.co.id", "Kalbe", 95, "A", 1),
        ])
        self.assertEqual(score_change_over_past_month("Kalbe.co.id", 90, self.path), 8.0)
        self.assertIsNone(score_change_over_past_month("other.com", 90, self.path))

    def test_unreadable_workbook_is_treated_as_missing_history(self):
        """A corrupt or half-synced workbook must not crash dashboard start-up."""
        self.path.write_bytes(b"this is not an xlsx file")
        self.assertEqual(current_month_history_status(self.path), (None, 0))
        self.assertEqual(score_trends_from_history(self.path), {})


if __name__ == "__main__":
    unittest.main()
