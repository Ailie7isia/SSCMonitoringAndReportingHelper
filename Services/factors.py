"""Track SecurityScorecard's ten risk-factor scores in the score-history workbook.

Factor scores live on their own "Factor History" sheet next to the score
history: one row per domain per snapshot, with a score column and an issue
count column for every factor. Live snapshots come from
GET /companies/{domain}/factors; missing past months are backfilled from
GET /companies/{domain}/history/factors/score, which SecurityScorecard serves
as monthly averages for roughly the last year.
"""

from __future__ import annotations

import logging
import time
from dataclasses import dataclass, field
from datetime import datetime, timezone
from functools import partial
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from models import Company
from Services.rate_limit import RequestPacer, StatusCallback
from Services.scores import FACTOR_SHEET, GRADE_FILLS, SCORE_HISTORY_PATH, UNREADABLE_WORKBOOK_ERRORS
from ssc_client import SecurityScorecardClient
from utils import normalize_domain

# API key -> label, in the order SecurityScorecard's Detailed report lists them.
FACTORS: dict[str, str] = {
    "application_security": "Application Security",
    "cubit_score": "Cubit Score",
    "dns_health": "DNS Health",
    "endpoint_security": "Endpoint Security",
    "hacker_chatter": "Hacker Chatter",
    "ip_reputation": "IP Reputation",
    "leaked_information": "Information Leak",
    "network_security": "Network Security",
    "patching_cadence": "Patching Cadence",
    "social_engineering": "Social Engineering",
}
# Column headings for tables too narrow for the full labels.
SHORT_LABELS: dict[str, str] = {
    "application_security": "App Sec",
    "cubit_score": "Cubit",
    "dns_health": "DNS",
    "endpoint_security": "Endpoint",
    "hacker_chatter": "Hacker Chatter",
    "ip_reputation": "IP Rep",
    "leaked_information": "Info Leak",
    "network_security": "Network",
    "patching_cadence": "Patching",
    "social_engineering": "Social Eng",
}

LIVE_SOURCE = "Live"
HISTORY_SOURCE = "SSC monthly history"
BASE_HEADERS = ["Retrieved (UTC)", "Domain", "Company", "Source", "Cycle"]
HEADERS = BASE_HEADERS + list(FACTORS.values()) + [f"{label} issues" for label in FACTORS.values()]
# Issue types counted the way the Detailed report counts them; positive and
# informational findings are not issues.
ISSUE_SEVERITIES = {"low", "medium", "high"}
# Past months to backfill from SecurityScorecard's monthly history (it serves about a year).
BACKFILL_MONTHS = 12


def grade_for_score(score: float | None) -> str:
    """SecurityScorecard's letter grade for a 0-100 score, or "?" without one."""
    if score is None:
        return "?"
    if score >= 90:
        return "A"
    if score >= 80:
        return "B"
    if score >= 70:
        return "C"
    if score >= 60:
        return "D"
    return "F"


@dataclass(frozen=True, slots=True)
class FactorSnapshot:
    """One domain's factor scores at one point in time."""

    retrieved: datetime  # naive UTC, like the score history
    domain: str
    company: str
    source: str  # LIVE_SOURCE or HISTORY_SOURCE
    cycle: int | str | None
    scores: dict[str, float | None]
    # Issue types per factor; None when the source does not report them.
    issues: dict[str, int | None] = field(default_factory=dict)

    @property
    def month(self) -> tuple[int, int]:
        return self.retrieved.year, self.retrieved.month


@dataclass(slots=True)
class FactorUpdate:
    """What one factor-history update saved."""

    live: int = 0
    backfilled: int = 0
    failed: list[str] = field(default_factory=list)


def _number(value: object) -> float | None:
    try:
        return float(value)  # type: ignore[arg-type]
    except (TypeError, ValueError):
        return None


def snapshot_from_factors(
    company: Company,
    entries: Iterable[dict],
    *,
    cycle: int | str | None,
    retrieved: datetime,
) -> FactorSnapshot | None:
    """Build a live snapshot from GET /companies/{domain}/factors, or None without scores."""
    scores: dict[str, float | None] = {}
    issues: dict[str, int | None] = {}
    for entry in entries:
        key = str(entry.get("name") or "").strip().lower()
        if key not in FACTORS:
            continue
        scores[key] = _number(entry.get("score"))
        issues[key] = sum(
            1
            for item in entry.get("issue_summary") or []
            if str(item.get("severity") or "").lower() in ISSUE_SEVERITIES
        )
    if all(score is None for score in scores.values()):
        return None
    return FactorSnapshot(
        retrieved=retrieved,
        domain=normalize_domain(company.domain),
        company=company.name,
        source=LIVE_SOURCE,
        cycle=cycle,
        scores=scores,
        issues=issues,
    )


def snapshots_from_history(company: Company, entries: Iterable[dict]) -> list[FactorSnapshot]:
    """Convert GET /companies/{domain}/history/factors/score entries into snapshots.

    Scores there are monthly averages such as 99.99999999999997, so they are
    rounded to one decimal place.
    """
    snapshots: list[FactorSnapshot] = []
    for entry in entries:
        try:
            retrieved = datetime.fromisoformat(str(entry.get("date") or "").replace("Z", "+00:00"))
        except ValueError:
            continue
        if retrieved.tzinfo is not None:
            retrieved = retrieved.astimezone(timezone.utc).replace(tzinfo=None)
        scores: dict[str, float | None] = {}
        for factor in entry.get("factors") or []:
            key = str(factor.get("name") or "").strip().lower()
            score = _number(factor.get("score"))
            if key in FACTORS:
                scores[key] = round(score, 1) if score is not None else None
        if any(score is not None for score in scores.values()):
            snapshots.append(
                FactorSnapshot(
                    retrieved=retrieved,
                    domain=normalize_domain(company.domain),
                    company=company.name,
                    source=HISTORY_SOURCE,
                    cycle=None,
                    scores=scores,
                )
            )
    return snapshots


def latest_by_month(snapshots: Iterable[FactorSnapshot]) -> dict[str, dict[tuple[int, int], FactorSnapshot]]:
    """Newest snapshot per domain per calendar month, matching the score history."""
    result: dict[str, dict[tuple[int, int], FactorSnapshot]] = {}
    for snapshot in snapshots:
        months = result.setdefault(snapshot.domain, {})
        current = months.get(snapshot.month)
        if current is None or snapshot.retrieved > current.retrieved:
            months[snapshot.month] = snapshot
    return result


def _prepare_sheet(worksheet) -> dict[str, int]:
    """Write the title and headers on a new sheet and add any missing columns.

    Returns each header's column number, so rows are written by name and
    sheets created before a factor existed keep working.
    """
    if worksheet["A1"].value is None:
        worksheet["A1"] = "SSC Monitoring Helper — Factor History"
        worksheet["A1"].font = Font(bold=True, size=16, color="FFFFFF")
        worksheet["A1"].fill = PatternFill("solid", fgColor="14212D")
        worksheet.freeze_panes = "C3"
    columns = {str(cell.value): cell.column for cell in worksheet[2] if cell.value is not None}
    for header in HEADERS:
        if header in columns:
            continue
        column = max(columns.values(), default=0) + 1
        cell = worksheet.cell(row=2, column=column, value=header)
        cell.font = Font(bold=True, color="10202A")
        cell.fill = PatternFill("solid", fgColor="2FBF8F")
        columns[header] = column
    return columns


def append_factor_history(
    snapshots: list[FactorSnapshot],
    output_file: Path = SCORE_HISTORY_PATH,
) -> int:
    """Append snapshots to the Factor History sheet; returns how many were written."""
    if not snapshots:
        return 0
    if output_file.exists():
        workbook = load_workbook(output_file)
    else:
        output_file.parent.mkdir(parents=True, exist_ok=True)
        workbook = Workbook()
        workbook.active.title = FACTOR_SHEET
    try:
        worksheet = workbook[FACTOR_SHEET] if FACTOR_SHEET in workbook.sheetnames else workbook.create_sheet(FACTOR_SHEET)
        columns = _prepare_sheet(worksheet)
        for snapshot in sorted(snapshots, key=lambda item: (item.domain, item.retrieved)):
            row = worksheet.max_row + 1
            values: dict[str, object] = {
                "Retrieved (UTC)": snapshot.retrieved,
                "Domain": snapshot.domain,
                "Company": snapshot.company,
                "Source": snapshot.source,
                "Cycle": snapshot.cycle,
            }
            for key, label in FACTORS.items():
                values[label] = snapshot.scores.get(key)
                values[f"{label} issues"] = snapshot.issues.get(key)
            for header, value in values.items():
                worksheet.cell(row=row, column=columns[header], value=value)
            worksheet.cell(row=row, column=columns["Retrieved (UTC)"]).number_format = "yyyy-mm-dd hh:mm:ss"
            for key, label in FACTORS.items():
                grade = grade_for_score(snapshot.scores.get(key))
                if grade in GRADE_FILLS:
                    worksheet.cell(row=row, column=columns[label]).fill = PatternFill("solid", fgColor=GRADE_FILLS[grade])
        try:
            workbook.save(output_file)
        except PermissionError as exc:
            raise PermissionError(
                f"Cannot update {output_file.name} because it is open in Excel. "
                "Close the workbook and try Update score again."
            ) from exc
    finally:
        workbook.close()
    return len(snapshots)


def read_factor_history(history_file: Path = SCORE_HISTORY_PATH) -> list[FactorSnapshot]:
    """Read every saved factor snapshot; a missing or unreadable workbook reads as none."""
    if not history_file.exists():
        return []
    try:
        workbook = load_workbook(history_file, read_only=True, data_only=True)
        try:
            if FACTOR_SHEET not in workbook.sheetnames:
                return []
            rows = workbook[FACTOR_SHEET].iter_rows(min_row=2, values_only=True)
            header = next(rows, None) or ()
            index = {str(value): position for position, value in enumerate(header) if value is not None}
            if not {"Retrieved (UTC)", "Domain"} <= index.keys():
                return []

            def value(row: tuple, name: str) -> object:
                position = index.get(name)
                return row[position] if position is not None and position < len(row) else None

            snapshots: list[FactorSnapshot] = []
            for row in rows:
                retrieved, domain = value(row, "Retrieved (UTC)"), value(row, "Domain")
                if not isinstance(retrieved, datetime) or not domain:
                    continue
                scores = {key: _number(value(row, label)) for key, label in FACTORS.items()}
                if all(score is None for score in scores.values()):
                    continue
                issues: dict[str, int | None] = {}
                for key, label in FACTORS.items():
                    count = _number(value(row, f"{label} issues"))
                    issues[key] = int(count) if count is not None else None
                snapshots.append(
                    FactorSnapshot(
                        retrieved=retrieved.replace(tzinfo=None),
                        domain=normalize_domain(str(domain)),
                        company=str(value(row, "Company") or domain).strip(),
                        source=str(value(row, "Source") or LIVE_SOURCE),
                        cycle=value(row, "Cycle"),  # type: ignore[arg-type]
                        scores=scores,
                        issues=issues,
                    )
                )
            return snapshots
        finally:
            workbook.close()
    except UNREADABLE_WORKBOOK_ERRORS:
        return []


def _past_months(now: datetime, count: int) -> set[tuple[int, int]]:
    """The ``count`` calendar months before ``now``'s month."""
    months: set[tuple[int, int]] = set()
    year, month = now.year, now.month
    for _ in range(count):
        month -= 1
        if month == 0:
            year, month = year - 1, 12
        months.add((year, month))
    return months


def update_factor_history(
    client: SecurityScorecardClient,
    companies: list[Company],
    cycle: int | str | None,
    output_file: Path = SCORE_HISTORY_PATH,
    *,
    backfill: bool = True,
    on_status: StatusCallback | None = None,
) -> FactorUpdate:
    """Save a live factor snapshot per company, backfilling missing past months first.

    Backfilling only calls the history endpoint for a domain that is missing
    one of the past BACKFILL_MONTHS months, so it runs once per new domain
    rather than on every update. The current month always comes from the live
    snapshot. A domain whose factors cannot be read is reported and skipped.
    """
    pacer = RequestPacer(
        # Looked up on each call so tests can patch the time module.
        sleep=lambda seconds: time.sleep(seconds),
        clock=lambda: time.monotonic(),
        on_status=on_status,
    )
    saved = latest_by_month(read_factor_history(output_file))
    now = datetime.now(timezone.utc).replace(tzinfo=None)
    wanted = _past_months(now, BACKFILL_MONTHS)
    result = FactorUpdate()
    snapshots: list[FactorSnapshot] = []
    try:
        for company in companies:
            domain = normalize_domain(company.domain)
            try:
                have = set(saved.get(domain, {}))
                if backfill and wanted - have:
                    entries = pacer.run(
                        "companies/history",
                        partial(client.get_factor_history, company.domain),
                        f"loading factor history for {domain}",
                    )
                    by_month = latest_by_month(snapshots_from_history(company, entries)).get(domain, {})
                    added = [snapshot for month, snapshot in by_month.items() if month in wanted - have]
                    snapshots.extend(added)
                    result.backfilled += len(added)
                entries = pacer.run(
                    "companies/factors",
                    partial(client.get_company_factors, company.domain),
                    f"loading factors for {domain}",
                )
                live = snapshot_from_factors(company, entries, cycle=cycle, retrieved=now)
                if live is None:
                    raise ValueError("SecurityScorecard returned no factor scores")
                snapshots.append(live)
                result.live += 1
            except Exception as exc:
                logging.warning("[%s] Factor scores unavailable: %s", domain, exc)
                result.failed.append(domain)
        append_factor_history(snapshots, output_file)
    finally:
        logging.info(pacer.finish())
    return result
