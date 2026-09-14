"""Windows dashboard for SecurityScorecard portfolio operations."""

from __future__ import annotations

import csv
import logging
import os
import queue
import threading
from datetime import datetime, timedelta, timezone
from pathlib import Path
from tkinter import messagebox
from typing import Callable

import customtkinter as ctk
from openpyxl import load_workbook

from config import CONFIG_PATH, REPORTS_DIR, load_config, validate_ssc_config
from constants import OPTION_LABELS, OPTION_VENDORS
from models import Company
from Services.portfolio import (
    apply_plan,
    companies_for_cycle_action,
    companies_from_payload,
    compute_plan,
    target_domains,
)
from Services.reports import download_reports
from Services.scores import (
    SCORE_HISTORY_PATH,
    append_score_history,
    current_month_history_status,
    grade_statistics,
    score_change_over_past_month,
    score_trends_from_history,
)
from ssc_client import ApiRequestError, RateLimitError, SecurityScorecardClient
from utils import format_duration, normalize_domain, sanitize_filename

ctk.set_appearance_mode("dark")
ctk.set_default_color_theme("green")

BACKGROUND, PANEL, PANEL_ALT = "#101923", "#182532", "#213240"
ACCENT, MUTED = "#2FBF8F", "#9CB2C2"
GRADE_COLORS = {"A": "#35C98A", "B": "#F1BE4D", "C": "#F48F4A", "D": "#EE5D69", "F": "#C93546"}
REPORT_GRADE_STYLES = {
    "A": ("A (90 – 100)", "#3FA548", "#0A2A0C"),
    "B": ("B (80 – 89)", "#F2B705", "#3D2B00"),
    "C": ("C (70 – 79)", "#F0790F", "#3D1900"),
    "D": ("D (60 – 69)", "#D5312E", "#FFFFFF"),
    "F": ("F (di bawah 60)", "#8B1A1A", "#FFFFFF"),
}


class QueueLogHandler(logging.Handler):
    def __init__(self, destination: "queue.Queue[str]") -> None:
        super().__init__()
        self.destination = destination

    def emit(self, record: logging.LogRecord) -> None:
        message = self.format(record)
        # Errors from the services are marked so the Activity panel shows them in red.
        if record.levelno >= logging.ERROR:
            message = f"✗ {message}"
        self.destination.put(message)


class Dashboard(ctk.CTk):
    def __init__(self) -> None:
        super().__init__()
        self.title("SSC Monitoring Helper")
        self.geometry("1240x790")
        self.minsize(1050, 680)
        self.configure(fg_color=BACKGROUND)
        self.log_queue: "queue.Queue[str]" = queue.Queue()
        self.busy = False
        self.current_companies: list[Company] = []
        self.active_cycle: int | None = None
        self.trends_page: ctk.CTkFrame | None = None
        self.report_page: ctk.CTkFrame | None = None
        self._rate_limit_resume_at: datetime | None = None
        self._rate_limit_job: str | None = None
        self._build_layout()
        self._build_loading_overlay()
        self.after(150, self._poll_log_queue)
        self.after(350, self.refresh_dashboard)

    def _build_layout(self) -> None:
        self.grid_columnconfigure(1, weight=1)
        self.grid_rowconfigure(1, weight=1)
        sidebar = ctk.CTkFrame(self, width=255, corner_radius=0, fg_color="#14212D")
        sidebar.grid(row=0, column=0, rowspan=2, sticky="nsew")
        sidebar.grid_propagate(False)
        ctk.CTkLabel(sidebar, text="SSC", font=("Segoe UI", 29, "bold"), text_color=ACCENT).pack(anchor="w", padx=27, pady=(31, 0))
        ctk.CTkLabel(sidebar, text="MONITORING HELPER", font=("Segoe UI", 14, "bold"), text_color=MUTED).pack(anchor="w", padx=29, pady=(0, 37))
        self.navigation_buttons: list[ctk.CTkButton] = []
        self._side_button(sidebar, "▣   Dashboard", self.show_dashboard)
        self._side_button(sidebar, "↗   Scores Trend", self.open_score_trends)
        self._side_button(sidebar, "▤   Report", self.open_report_page)
        self.sidebar_status = ctk.CTkLabel(
            sidebar,
            text="●  Ready",
            justify="left",
            wraplength=190,
            font=("Segoe UI", 12, "bold"),
            text_color=ACCENT,
        )
        self.sidebar_status.pack(side="bottom", anchor="w", fill="x", padx=24, pady=22)
        # Shown above the status only while SecurityScorecard's request quota is exhausted.
        self.rate_limit_status = ctk.CTkLabel(
            sidebar,
            text="",
            justify="left",
            wraplength=200,
            font=("Segoe UI", 12, "bold"),
            text_color="#F2B84B",
        )

        self.dashboard_header = ctk.CTkFrame(self, height=112, corner_radius=0, fg_color=BACKGROUND)
        self.dashboard_header.grid(row=0, column=1, sticky="new", padx=35)
        self.dashboard_header.grid_columnconfigure(0, weight=1)
        ctk.CTkLabel(self.dashboard_header, text="Dashboard", font=("Segoe UI", 29, "bold")).grid(row=0, column=0, sticky="sw", pady=(25, 0))
        self.last_updated = ctk.CTkLabel(self.dashboard_header, text="Loading portfolio data…", font=("Segoe UI", 14), text_color=MUTED)
        self.last_updated.grid(row=1, column=0, sticky="nw", pady=(2, 0))
        self.refresh_button = ctk.CTkButton(self.dashboard_header, text="Refresh now", width=125, height=34, command=self.refresh_dashboard, fg_color=PANEL_ALT, hover_color="#2C4354", font=("Segoe UI", 14))
        self.refresh_button.grid(row=0, column=1, rowspan=2, sticky="e")

        self.dashboard_body = ctk.CTkScrollableFrame(self, fg_color=BACKGROUND, corner_radius=0)
        self.dashboard_body.grid(row=1, column=1, sticky="nsew", padx=(31, 23), pady=(0, 20))
        body = self.dashboard_body
        body.grid_columnconfigure((0, 1, 2), weight=1)
        self.metric_values: dict[str, ctk.CTkLabel] = {}
        self.metric_subtitles: dict[str, ctk.CTkLabel] = {}
        for index, (key, title, value, subtitle, color) in enumerate((
            ("companies", "Active domains", "—", "Current portfolio", ACCENT),
            ("kalbe", "Kalbe.co.id score", "—", "No 30-day comparison yet", "#6BAFFF"),
            ("attention", "Needs attention", "—", "Grade B and below", "#F2B84B"),
        )):
            card = self._metric_card(body, title, value, subtitle, color)
            card.grid(row=0, column=index, sticky="ew", padx=7, pady=(5, 14))
            self.metric_values[key] = card.value_label  # type: ignore[attr-defined]
            self.metric_subtitles[key] = card.subtitle_label  # type: ignore[attr-defined]

        domains = ctk.CTkFrame(body, fg_color=PANEL, corner_radius=14)
        domains.grid(row=1, column=0, columnspan=2, sticky="nsew", padx=7, pady=7)
        domains.grid_columnconfigure(0, weight=1)
        ctk.CTkLabel(domains, text="Active Domains", font=("Segoe UI", 17, "bold")).grid(row=0, column=0, sticky="w", padx=20, pady=(17, 1))
        self.cycle_context = ctk.CTkLabel(
            domains,
            text="Currently viewing active portfolio",
            font=("Segoe UI", 14),
            text_color=MUTED,
        )
        self.cycle_context.grid(row=1, column=0, sticky="w", padx=20)
        self.copy_domains_button = ctk.CTkButton(
            domains,
            text="Copy all",
            width=98,
            height=30,
            command=self._copy_active_domains,
            fg_color=PANEL_ALT,
            hover_color="#2C4354",
            font=("Segoe UI", 13, "bold"),
            state="disabled",
        )
        self.copy_domains_button.grid(row=0, column=1, rowspan=2, sticky="e", padx=20, pady=(12, 0))
        self.domain_rows = ctk.CTkScrollableFrame(domains, height=245, fg_color="#0E1821", corner_radius=9)
        self.domain_rows.grid(row=2, column=0, columnspan=2, sticky="ew", padx=20, pady=(12, 20))
        self.domain_rows.grid_columnconfigure(0, weight=1)
        self._render_domains([])

        actions = ctk.CTkFrame(body, fg_color=PANEL, corner_radius=14)
        actions.grid(row=1, column=2, sticky="nsew", padx=7, pady=7)
        self.action_buttons: list[ctk.CTkButton] = []
        self._action_button(actions, "◈   Run portfolio cycle", self.open_cycle_dialog, primary=True)
        self._action_button(actions, "↓   Download detailed reports", self.start_detailed_reports_download)
        self._action_button(actions, "↓   Download issue reports", self.start_issue_reports_download)
        self._action_button(actions, "▣   Open reports folder", self.open_reports_folder)
        self._action_button(actions, "↑   Update score history", self.start_score_export)
        self.history_status = ctk.CTkLabel(
            actions,
            text="Checking score history…",
            justify="left",
            wraplength=240,
            font=("Segoe UI", 12),
            text_color=MUTED,
        )
        self.history_status.pack(anchor="w", padx=20, pady=(7, 14))
        self._update_history_status()

        activity = ctk.CTkFrame(body, fg_color=PANEL, corner_radius=14)
        activity.grid(row=2, column=0, columnspan=3, sticky="nsew", padx=7, pady=(15, 8))
        activity.grid_columnconfigure(0, weight=1)
        ctk.CTkLabel(activity, text="Activity", font=("Segoe UI", 17, "bold")).grid(row=0, column=0, sticky="w", padx=20, pady=(17, 7))
        self.activity = ctk.CTkTextbox(activity, height=188, font=("Cascadia Mono", 14), fg_color="#0E1821", border_width=0)
        self.activity.grid(row=1, column=0, sticky="nsew", padx=20, pady=(0, 20))
        self.activity.tag_config("failure", foreground="#FF6B73")
        self.activity.insert("end", "Monitoring Helper started. Loading current portfolio…\n")
        self.activity.configure(state="disabled")

    def _build_loading_overlay(self) -> None:
        """A modal, full-window state keeps long API operations unambiguous."""
        self.loading_overlay = ctk.CTkFrame(self, fg_color="#0B141D", corner_radius=0)
        panel = ctk.CTkFrame(self.loading_overlay, width=520, height=250, fg_color=PANEL, corner_radius=18)
        panel.place(relx=0.5, rely=0.5, anchor="center")
        panel.pack_propagate(False)
        ctk.CTkLabel(panel, text="SECURITYSCORECARD", font=("Segoe UI", 14, "bold"), text_color=ACCENT).pack(pady=(45, 8))
        self.loading_title = ctk.CTkLabel(panel, text="Loading portfolio", font=("Segoe UI", 25, "bold"))
        self.loading_title.pack(pady=(0, 8))
        ctk.CTkLabel(panel, text="Please wait — Monitoring Helper will update automatically.", font=("Segoe UI", 14), text_color=MUTED).pack(pady=(0, 25))
        self.loading_bar = ctk.CTkProgressBar(panel, width=390, height=14, mode="indeterminate", progress_color=ACCENT)
        self.loading_bar.pack()

    def _side_button(self, parent: ctk.CTkFrame, text: str, command: Callable[[], None]) -> None:
        button = ctk.CTkButton(parent, text=text, command=command, anchor="w", height=43, corner_radius=8, fg_color="transparent", hover_color=PANEL_ALT, font=("Segoe UI", 14))
        button.pack(fill="x", padx=17, pady=(3, 0))
        self.navigation_buttons.append(button)

    def _action_button(
        self,
        parent: ctk.CTkFrame,
        text: str,
        command: Callable[[], None],
        *,
        primary: bool = False,
    ) -> None:
        button = ctk.CTkButton(
            parent,
            text=text,
            command=command,
            anchor="w",
            height=40,
            corner_radius=8,
            fg_color=ACCENT if primary else PANEL_ALT,
            hover_color="#249E74" if primary else "#2C4354",
            font=("Segoe UI", 14),
        )
        button.pack(fill="x", padx=20, pady=(18, 0) if primary else (8, 0))
        self.action_buttons.append(button)

    @staticmethod
    def _metric_card(parent: ctk.CTkFrame, title: str, value: str, subtitle: str, color: str) -> ctk.CTkFrame:
        card = ctk.CTkFrame(parent, fg_color=PANEL, corner_radius=14)
        ctk.CTkLabel(card, text=title.upper(), font=("Segoe UI", 14, "bold"), text_color=MUTED).pack(anchor="w", padx=19, pady=(16, 0))
        label = ctk.CTkLabel(card, text=value, font=("Segoe UI", 30, "bold"), text_color=color)
        label.pack(anchor="w", padx=19, pady=(1, 0))
        subtitle_label = ctk.CTkLabel(card, text=subtitle, font=("Segoe UI", 14), text_color=MUTED)
        subtitle_label.pack(anchor="w", padx=19, pady=(0, 15))
        card.value_label = label  # type: ignore[attr-defined]
        card.subtitle_label = subtitle_label  # type: ignore[attr-defined]
        return card

    def _set_busy(
        self,
        busy: bool,
        status: str = "Ready",
        *,
        show_loading_overlay: bool = True,
    ) -> None:
        self.busy = busy
        state = "disabled" if busy else "normal"
        for button in self.navigation_buttons:
            button.configure(state=state)
        self.refresh_button.configure(state=state)
        for button in self.action_buttons:
            button.configure(state=state)
        self.sidebar_status.configure(text=f"{'◌' if busy else '●'}  {status}", text_color="#F2B84B" if busy else ACCENT)
        if busy and show_loading_overlay:
            self.loading_title.configure(text=status)
            self.loading_overlay.place(relx=0, rely=0, relwidth=1, relheight=1)
            self.loading_bar.start()
        else:
            self.loading_bar.stop()
            self.loading_overlay.place_forget()

    def _log(self, text: str) -> None:
        self.log_queue.put(text)

    def _show_rate_limit_countdown(self, resume_at: datetime | None) -> None:
        """Show when SecurityScorecard's request quota is expected to reset."""
        if self._rate_limit_job is not None:
            self.after_cancel(self._rate_limit_job)
            self._rate_limit_job = None
        self._rate_limit_resume_at = resume_at
        self._tick_rate_limit_countdown()

    def _tick_rate_limit_countdown(self) -> None:
        resume_at = self._rate_limit_resume_at
        remaining = (resume_at - datetime.now().astimezone()).total_seconds() if resume_at else 0
        if remaining <= 0:
            self._rate_limit_resume_at = None
            self._rate_limit_job = None
            self.rate_limit_status.pack_forget()
            return
        self.rate_limit_status.configure(
            text=f"⏳  SSC request limit reached\nResumes in {format_duration(remaining)} (at {resume_at:%H:%M:%S})"
        )
        self.rate_limit_status.pack(side="bottom", anchor="w", fill="x", padx=24, pady=(0, 2))
        self._rate_limit_job = self.after(1000, self._tick_rate_limit_countdown)

    def _report_rate_limit(self, exc: RateLimitError) -> None:
        """Log a rate-limited operation and start the countdown when the reset time is known."""
        if exc.retry_after is None:
            self._log(
                "✗ SecurityScorecard request limit reached. It did not say when the quota resets; "
                "its general limit is a rolling 60-minute window, so try again later."
            )
            return
        resume_at = datetime.now().astimezone() + timedelta(seconds=exc.retry_after)
        self._log(
            f"✗ SecurityScorecard request limit reached. Try again after {resume_at:%H:%M:%S} "
            f"(in {format_duration(exc.retry_after)})."
        )
        self.after(0, lambda: self._show_rate_limit_countdown(resume_at))

    def _update_history_status(self) -> tuple[datetime | None, int]:
        """Show the most recent score-history update made this month."""
        saved_at, records = current_month_history_status()
        if saved_at is not None:
            # The workbook stores UTC; show the operator's local time.
            saved_at = saved_at.replace(tzinfo=timezone.utc).astimezone()
        if saved_at is None:
            self.history_status.configure(
                text="○  No score history saved\nthis month yet",
                text_color="#F2B84B",
            )
        else:
            self.history_status.configure(
                text=f"✓  Latest history update\n{saved_at:%d %b %Y, %H:%M} • {records} records this month",
                text_color=ACCENT,
            )
        return saved_at, records

    def _poll_log_queue(self) -> None:
        try:
            while True:
                line = self.log_queue.get_nowait()
                self.activity.configure(state="normal")
                message = line.rstrip() + "\n"
                is_failure = message.lstrip().startswith("✗")
                self.activity.insert("end", message, "failure" if is_failure else None)
                self.activity.see("end")
                self.activity.configure(state="disabled")
        except queue.Empty:
            pass
        self.after(150, self._poll_log_queue)

    @staticmethod
    def _client() -> tuple[SecurityScorecardClient, str]:
        config = load_config(CONFIG_PATH)
        api_key, portfolio_id = validate_ssc_config(config)
        return SecurityScorecardClient(api_key), portfolio_id

    def _run(
        self,
        label: str,
        task: Callable[[], None],
        *,
        show_loading_overlay: bool = True,
    ) -> None:
        if self.busy:
            return
        self._set_busy(True, label, show_loading_overlay=show_loading_overlay)
        self._log(f"› {label}")
        def worker() -> None:
            try:
                task()
            except RateLimitError as exc:
                self._report_rate_limit(exc)
            except ApiRequestError as exc:
                self._log(f"✗ SecurityScorecard request failed: {exc}")
            except (FileNotFoundError, ValueError) as exc:
                self._log(f"✗ Configuration issue: {exc}")
            except Exception as exc:
                self._log(f"✗ Operation failed: {exc}")
            finally:
                self.after(0, lambda: self._set_busy(False))
        threading.Thread(target=worker, daemon=True).start()

    def refresh_dashboard(self) -> None:
        def load() -> None:
            client, portfolio_id = self._client()
            companies = companies_from_payload(client.fetch_portfolio_companies(portfolio_id))
            enriched: list[Company] = []
            for company in companies:
                details = client.get_company(company.domain)
                enriched.append(Company(company.domain, company.name, str(details.get("grade") or "").upper(), details.get("score")))
            kalbe = next((company for company in enriched if company.domain.lower() == "kalbe.co.id"), None)
            kalbe_trend = score_change_over_past_month(kalbe.domain, kalbe.score) if kalbe else None
            self.after(0, lambda: self._update_dashboard(enriched, kalbe_trend))
            self._log(f"✓ Portfolio refreshed: {len(enriched)} companies loaded.")
        self._run("Refreshing portfolio", load)

    def _update_dashboard(
        self,
        companies: list[Company],
        kalbe_trend: float | None,
    ) -> None:
        self.current_companies = companies
        self.cycle_context.configure(text=f"Currently viewing {self._current_cycle_label(companies)}")
        grades = grade_statistics(companies)
        attention = sum(grades.get(grade, 0) for grade in ("B", "C", "D", "F"))
        kalbe = next((company for company in companies if company.domain.lower() == "kalbe.co.id"), None)
        self.metric_values["companies"].configure(text=str(len(companies)))
        if kalbe and isinstance(kalbe.score, (int, float)):
            grade = kalbe.grade.upper() if kalbe.grade else "Unknown"
            self.metric_values["kalbe"].configure(
                text=f"{kalbe.score:.0f}",
                text_color=GRADE_COLORS.get(grade, MUTED),
            )
            self._set_kalbe_trend(kalbe_trend)
        else:
            self.metric_values["kalbe"].configure(text="—", text_color=MUTED)
            self.metric_subtitles["kalbe"].configure(text="Kalbe.co.id is not in the active portfolio", text_color=MUTED)
        self.metric_values["attention"].configure(text=str(attention))
        self.last_updated.configure(text=f"Last refreshed {datetime.now().strftime('%d %b %Y, %H:%M')}  •  {len(companies)} domains active")
        self._render_domains(companies)

    def _current_cycle_label(self, companies: list[Company]) -> str:
        """Return the configured cycle name when the portfolio matches one."""
        self.active_cycle = None
        current_domains = {company.domain.lower() for company in companies}
        for option in sorted(OPTION_VENDORS):
            if current_domains == {domain.lower() for domain in target_domains(option)}:
                self.active_cycle = option
                break
        if self.active_cycle is None:
            return "active portfolio"
        return f"Cycle {self.active_cycle} — {OPTION_LABELS[self.active_cycle]}"

    def _set_kalbe_trend(self, trend: float | None) -> None:
        if trend is None:
            self.metric_subtitles["kalbe"].configure(
                text="No 30-day comparison yet",
                text_color=MUTED,
            )
            return
        amount = f"{abs(trend):.0f}"
        if trend > 0:
            text, color = f"↑ +{amount} vs. 30 days ago", ACCENT
        elif trend < 0:
            text, color = f"↓ −{amount} vs. 30 days ago", "#EE5D69"
        else:
            text, color = "→ No change vs. 30 days ago", MUTED
        self.metric_subtitles["kalbe"].configure(text=text, text_color=color)

    def _render_domains(self, companies: list[Company]) -> None:
        for widget in self.domain_rows.winfo_children():
            widget.destroy()
        if not companies:
            self.copy_domains_button.configure(state="disabled")
            ctk.CTkLabel(self.domain_rows, text="No active domains loaded yet.", font=("Segoe UI", 14), text_color=MUTED).grid(row=0, column=0, sticky="w", padx=12, pady=12)
            return
        self.copy_domains_button.configure(state="normal")
        for row, company in enumerate(sorted(companies, key=lambda item: item.domain.lower())):
            grade = company.grade.upper() if company.grade else "Unknown"
            color = GRADE_COLORS.get(grade, "#718899")
            line = ctk.CTkFrame(self.domain_rows, fg_color="transparent")
            line.grid(row=row, column=0, sticky="ew", padx=5, pady=2)
            line.grid_columnconfigure(0, weight=1)
            ctk.CTkButton(
                line,
                text=company.domain,
                anchor="w",
                command=lambda selected=company: self.open_domain_details(selected),
                fg_color="transparent",
                hover_color=PANEL_ALT,
                font=("Segoe UI", 14),
            ).grid(row=0, column=0, sticky="ew", padx=(3, 12), pady=2)
            score = "—" if company.score is None else str(company.score)
            ctk.CTkLabel(line, text=score, width=48, corner_radius=7, fg_color=color, text_color="#10202A", font=("Segoe UI", 14, "bold")).grid(row=0, column=1, padx=(0, 7), pady=4)
            ctk.CTkLabel(line, text=grade, width=58, text_color=color, font=("Segoe UI", 14, "bold")).grid(row=0, column=2, padx=(0, 7), pady=4)

    def _copy_active_domains(self) -> None:
        """Copy every currently active domain and its latest score details."""
        rows = []
        for company in sorted(self.current_companies, key=lambda item: item.domain.lower()):
            grade = company.grade.upper() if company.grade else "Unknown"
            score = "—" if company.score is None else str(company.score)
            rows.append(f"{company.domain}\t{score}\t{grade}")
        if rows:
            self.clipboard_clear()
            self.clipboard_append("\n".join(rows))
            self._log(f"✓ Copied {len(rows)} active domains to the clipboard.")

    def open_cycle_dialog(self) -> None:
        if self.busy:
            return
        dialog = ctk.CTkToplevel(self)
        dialog.title("Run portfolio cycle")
        dialog.geometry("470x610")
        dialog.resizable(False, False)
        dialog.grab_set()
        ctk.CTkLabel(dialog, text="Choose a portfolio cycle", font=("Segoe UI", 22, "bold")).pack(anchor="w", padx=28, pady=(28, 3))
        ctk.CTkLabel(dialog, text="The planned additions and removals will be shown before any changes are made.", justify="left", wraplength=390, font=("Segoe UI", 14), text_color=MUTED).pack(anchor="w", padx=28, pady=(0, 18))
        cycle_options = sorted(OPTION_VENDORS)
        selected = ctk.IntVar(value=cycle_options[0])
        for option in cycle_options:
            cycle_name = OPTION_LABELS.get(option, f"Cycle {option}")
            ctk.CTkRadioButton(
                dialog,
                text=f"Cycle {option} — {cycle_name}",
                variable=selected,
                value=option,
                font=("Segoe UI", 14),
            ).pack(anchor="w", padx=31, pady=6)
        def continue_cycle() -> None:
            dialog.destroy()
            self._run_cycle(selected.get())
        ctk.CTkButton(dialog, text="Review and continue", command=continue_cycle, height=40, fg_color=ACCENT, hover_color="#249E74", font=("Segoe UI", 14)).pack(fill="x", padx=28, pady=27)

    def _run_cycle(self, option: int) -> None:
        def cycle() -> None:
            client, portfolio_id = self._client()
            plan = compute_plan(companies_from_payload(client.fetch_portfolio_companies(portfolio_id)), target_domains(option))
            summary = self._plan_text(option, plan.to_add, plan.to_remove)
            if not self._ask_confirmation("Confirm portfolio changes", summary + "\n\nApply these changes?"):
                self._log("• Portfolio cycle cancelled; no changes were made.")
                return
            report = apply_plan(client, portfolio_id, plan)
            self.active_cycle = option
            self._log(f"✓ Portfolio cycle complete: {len(report.added_ok)} added, {len(report.removed_ok)} removed, {len(report.added_failed) + len(report.removed_failed)} failed.")
            # Let the current task release the UI first, then reload the
            # management metrics from the updated portfolio.
            self.after(0, lambda: self.after(100, self.refresh_dashboard))
        self._run(f"Reviewing cycle {option}", cycle)

    @staticmethod
    def _plan_text(option: int, to_add: list[str], to_remove: list[str]) -> str:
        additions = "\n".join(f"  + {domain}" for domain in to_add) or "  None"
        removals = "\n".join(f"  − {domain}" for domain in to_remove) or "  None"
        return f"Cycle {option} portfolio plan\n\nAdd ({len(to_add)}):\n{additions}\n\nRemove ({len(to_remove)}):\n{removals}"

    def _ask_confirmation(self, title: str, message: str) -> bool:
        completed = threading.Event()
        answer: dict[str, bool] = {"value": False}

        def ask() -> None:
            # A native confirmation dialog should be the only modal surface the
            # operator sees. Hide the loading overlay until it is answered.
            self.loading_bar.stop()
            self.loading_overlay.place_forget()
            answer["value"] = messagebox.askyesno(title, message, parent=self)
            if answer["value"] and self.busy:
                self.loading_title.configure(text="Applying portfolio changes")
                self.loading_overlay.place(relx=0, rely=0, relwidth=1, relheight=1)
                self.loading_bar.start()
            completed.set()

        self.after(0, ask)
        completed.wait()
        return answer["value"]

    def start_detailed_reports_download(self) -> None:
        self._start_reports_download("detailed_report", "Detailed PDF reports")

    def start_issue_reports_download(self) -> None:
        self._start_reports_download("issue_report", "Issues CSV reports")

    def _start_reports_download(self, report_type: str, report_label: str) -> None:
        def download() -> None:
            client, portfolio_id = self._client()
            companies = companies_from_payload(client.fetch_portfolio_companies(portfolio_id))
            companies = companies_for_cycle_action(companies, self.active_cycle)
            if not companies:
                self._log("• No domains require reports for the active cycle.")
                return
            output_dir = REPORTS_DIR / datetime.now(timezone.utc).strftime("%Y-%m-%d")
            saved = download_reports(
                client,
                companies,
                output_dir,
                (report_type,),
                on_rate_limit=lambda resume_at: self.after(0, lambda: self._show_rate_limit_countdown(resume_at)),
            )
            self._log(f"✓ Downloaded {len(saved)} of {len(companies)} {report_label} to {output_dir}.")
        # Report generation can take a while. The service logs each request,
        # wait cycle, and saved file, so keep the Activity panel exposed.
        self._run(
            f"Generating and downloading {report_label.lower()}",
            download,
            show_loading_overlay=False,
        )

    def open_reports_folder(self) -> None:
        """Open the local report archive in Windows Explorer."""
        REPORTS_DIR.mkdir(parents=True, exist_ok=True)
        try:
            os.startfile(str(REPORTS_DIR))
        except OSError as exc:
            self._log(f"✗ Could not open the reports folder: {exc}")

    def _latest_issue_report(self, company: Company) -> Path | None:
        """Find the newest downloaded SecurityScorecard Issues CSV for a domain.

        Companies sharing a display name are saved with the domain appended,
        so that form is preferred over the plain company name.
        """
        markers = (
            f" - {sanitize_filename(f'{company.name} ({company.domain})')} - Issue Report - ",
            f" - {sanitize_filename(company.name)} - Issue Report - ",
        )
        dated_folders = sorted(
            (folder for folder in REPORTS_DIR.iterdir() if folder.is_dir()),
            key=lambda folder: folder.name,
            reverse=True,
        ) if REPORTS_DIR.exists() else []
        for folder in dated_folders:
            for marker in markers:
                matches = sorted(
                    (path for path in (folder / "issues").glob("*.csv") if marker in path.name),
                    key=lambda path: path.stat().st_mtime,
                    reverse=True,
                )
                if matches:
                    return matches[0]
        return None

    def open_domain_details(self, company: Company) -> None:
        """Display SecurityScorecard's saved findings and recommendations."""
        issue_report = self._latest_issue_report(company)
        if issue_report is None:
            self._log(f"• No downloaded Issues report is available for {company.domain}. Generate reports first.")
            return
        try:
            with open(issue_report, "r", encoding="utf-8-sig", newline="") as file:
                issues = list(csv.DictReader(file))
        except (OSError, csv.Error) as exc:
            self._log(f"✗ Could not read the Issues report for {company.domain}: {exc}")
            return
        dialog = ctk.CTkToplevel(self)
        dialog.title(f"{company.domain} — SecurityScorecard issues")
        dialog.geometry("930x690")
        dialog.minsize(720, 500)
        dialog.configure(fg_color=BACKGROUND)
        dialog.grid_columnconfigure(0, weight=1)
        dialog.grid_rowconfigure(2, weight=1)
        ctk.CTkLabel(dialog, text=company.domain, font=("Segoe UI", 25, "bold")).grid(
            row=0, column=0, sticky="w", padx=28, pady=(27, 0)
        )
        ctk.CTkLabel(
            dialog,
            text=f"{len(issues)} finding(s) from {issue_report.parent.parent.name} • SecurityScorecard Issues report",
            font=("Segoe UI", 14),
            text_color=MUTED,
        ).grid(row=1, column=0, sticky="w", padx=28, pady=(2, 14))
        findings = ctk.CTkScrollableFrame(dialog, fg_color=BACKGROUND, corner_radius=0)
        findings.grid(row=2, column=0, sticky="nsew", padx=20, pady=(0, 20))
        findings.grid_columnconfigure(0, weight=1)
        if not issues:
            ctk.CTkLabel(findings, text="SecurityScorecard reported no issues in this download.", font=("Segoe UI", 14), text_color=MUTED).grid(
                row=0, column=0, sticky="w", padx=12, pady=16
            )
            return
        for row, issue in enumerate(issues):
            card = ctk.CTkFrame(findings, fg_color=PANEL, corner_radius=12)
            card.grid(row=row, column=0, sticky="ew", padx=8, pady=7)
            card.grid_columnconfigure(0, weight=1)
            severity = (issue.get("ISSUE TYPE SEVERITY") or "Unknown").upper()
            title = issue.get("ISSUE TYPE TITLE") or issue.get("ISSUE TYPE CODE") or "Unnamed issue"
            ctk.CTkLabel(card, text=title, anchor="w", font=("Segoe UI", 15, "bold"), wraplength=700).grid(
                row=0, column=0, sticky="ew", padx=16, pady=(14, 2)
            )
            ctk.CTkLabel(
                card,
                text=f"{severity}  •  {issue.get('FACTOR NAME') or 'Uncategorized'}  •  {issue.get('STATUS') or 'Active'}",
                anchor="w",
                font=("Segoe UI", 14, "bold"),
                text_color=GRADE_COLORS.get("F" if severity in {"HIGH", "CRITICAL"} else "C", "#F2B84B"),
            ).grid(row=1, column=0, sticky="ew", padx=16, pady=(0, 8))
            target = issue.get("TARGET") or issue.get("HOSTNAME") or issue.get("IP ADDRESSES")
            if target:
                ctk.CTkLabel(card, text=f"Affected: {target}", anchor="w", font=("Segoe UI", 14), text_color=MUTED, wraplength=700).grid(
                    row=2, column=0, sticky="ew", padx=16, pady=(0, 8)
                )
            recommendation = issue.get("ISSUE RECOMMENDATION") or "No recommendation was included in this report."
            ctk.CTkLabel(card, text=f"Recommended solution\n{recommendation}", anchor="w", justify="left", font=("Segoe UI", 14), wraplength=700).grid(
                row=3, column=0, sticky="ew", padx=16, pady=(0, 15)
            )

    def open_score_trends(self) -> None:
        """Load score history without freezing the interface."""
        if not self.current_companies:
            self._log("• Refresh the portfolio before viewing score trends.")
            return
        if self.busy:
            return
        self._set_busy(True, "Loading score trends")

        def load() -> None:
            try:
                history_trends = score_trends_from_history()
                self.after(0, lambda: self._show_score_trends(history_trends))
            except Exception as exc:
                self.after(0, lambda: self._log(f"✗ Could not load score trends: {exc}"))
            finally:
                self.after(0, lambda: self._set_busy(False))

        threading.Thread(target=load, daemon=True).start()

    def _show_score_trends(self, history_trends: dict[str, object]) -> None:
        """Display the loaded score history in the trend page."""
        if self.report_page is not None:
            self.report_page.destroy()
            self.report_page = None
        if self.trends_page is not None:
            self.trends_page.destroy()
        self.dashboard_header.grid_remove()
        self.dashboard_body.grid_remove()

        self.trends_page = ctk.CTkFrame(self, fg_color=BACKGROUND, corner_radius=0)
        self.trends_page.grid(row=0, column=1, rowspan=2, sticky="nsew", padx=(31, 23), pady=(0, 20))
        self.trends_page.grid_columnconfigure(0, weight=1)
        # Keep the descriptive text compact so the results container receives
        # the available vertical space instead of leaving a large gap above it.
        self.trends_page.grid_rowconfigure(3, weight=1, minsize=540)
        ctk.CTkButton(
            self.trends_page, text="←  Dashboard", command=self.close_score_trends,
            width=170, height=32, fg_color="transparent", hover_color=PANEL_ALT, anchor="w", font=("Segoe UI", 14),
        ).grid(row=0, column=0, sticky="w", padx=4, pady=(22, 0))
        ctk.CTkLabel(self.trends_page, text="Scores Trend", font=("Segoe UI", 25, "bold")).grid(
            row=1, column=0, sticky="w", padx=4, pady=(13, 0)
        )
        ctk.CTkLabel(
            self.trends_page,
            text="Scores use the newest saved entry in each calendar month. Missing current-month updates display as —.",
            font=("Segoe UI", 14),
            text_color=MUTED,
        ).grid(row=2, column=0, sticky="nw", padx=4, pady=(2, 14))

        content = ctk.CTkFrame(self.trends_page, height=540, fg_color=PANEL, corner_radius=14)
        content.grid(row=3, column=0, sticky="nsew", padx=4, pady=(0, 8))
        content.grid_columnconfigure(0, weight=1)
        mode = ctk.StringVar(value="Active")
        search_text = ctk.StringVar()
        grade_filter = ctk.StringVar(value="All grades")
        switcher = ctk.CTkSegmentedButton(
            content,
            values=["Active", "All"],
            variable=mode,
            width=330,
            font=("Segoe UI", 14),
        )
        switcher.grid(row=0, column=0, sticky="w", padx=20, pady=(18, 8))
        filters = ctk.CTkFrame(content, fg_color="transparent")
        filters.grid(row=1, column=0, sticky="ew", padx=20, pady=(0, 12))
        filters.grid_columnconfigure(0, weight=1)
        search = ctk.CTkEntry(
            filters,
            textvariable=search_text,
            placeholder_text="Search domains…",
            height=36,
            font=("Segoe UI", 14),
        )
        search.grid(row=0, column=0, sticky="ew", padx=(0, 12))
        grade_menu = ctk.CTkOptionMenu(
            filters,
            variable=grade_filter,
            values=["All grades", "A (90–100)", "B (80–89)", "C (70–79)", "D (60–69)", "F (0–59)", "No score"],
            height=36,
            font=("Segoe UI", 14),
            dropdown_font=("Segoe UI", 14),
        )
        grade_menu.grid(row=0, column=1, sticky="e")
        table = ctk.CTkScrollableFrame(content, fg_color="#0E1821", corner_radius=9)
        table.grid(row=2, column=0, sticky="nsew", padx=20, pady=(0, 20))
        content.grid_rowconfigure(2, weight=1)
        table.grid_columnconfigure(0, weight=1)

        def delta(value: float | None) -> tuple[str, str]:
            if value is None:
                return "—", MUTED
            if value > 0:
                return f"↑ +{value:.0f}", ACCENT
            if value < 0:
                return f"↓ −{abs(value):.0f}", "#EE5D69"
            return "→ 0", MUTED

        def score_grade(score: float | None) -> str:
            if score is None:
                return "No score"
            if score >= 90:
                return "A"
            if score >= 80:
                return "B"
            if score >= 70:
                return "C"
            if score >= 60:
                return "D"
            return "F"

        def render(_selection: str = "") -> None:
            for widget in table.winfo_children():
                widget.destroy()
            headers = (("Domain", "w"), ("This month score", "e"), ("Last month score", "e"), ("MoM", "e"), ("YoY", "e"))
            for column, (label, anchor) in enumerate(headers):
                table.grid_columnconfigure(column, weight=1 if column == 0 else 0)
                ctk.CTkLabel(
                    table,
                    text=label.upper(),
                    anchor=anchor,
                    font=("Segoe UI", 16 if label == "MoM" else 14, "bold"),
                    text_color=ACCENT if label == "MoM" else MUTED,
                ).grid(
                    row=0, column=column, sticky="ew", padx=12, pady=(10, 7)
                )
            if mode.get() == "Active":
                rows = [
                    (company.domain, history_trends.get(company.domain.lower()))
                    for company in self.current_companies
                ]
            else:
                rows = [
                    (domain, trend)
                    for domain, trend in history_trends.items()
                ]
            query = search_text.get().strip().lower()
            selected_grade = "No score" if grade_filter.get() == "No score" else grade_filter.get().split(" ", 1)[0]
            rows = [
                (domain, trend)
                for domain, trend in rows
                if (not query or query in domain.lower())
                and (selected_grade == "All" or score_grade(trend.this_month_score if trend else None) == selected_grade)
            ]
            if not rows:
                ctk.CTkLabel(table, text="No domains match the current search and grade filter.", font=("Segoe UI", 14), text_color=MUTED).grid(
                    row=1, column=0, columnspan=5, sticky="w", padx=12, pady=12
                )
                return
            for row, (domain, trend) in enumerate(sorted(rows, key=lambda item: item[0].lower()), start=1):
                month_text, month_color = delta(trend.month_over_month if trend else None)
                year_text, year_color = delta(trend.year_over_year if trend else None)
                this_month_score = "—" if trend is None or trend.this_month_score is None else f"{trend.this_month_score:.0f}"
                last_month_score = "—" if trend is None or trend.last_month_score is None else f"{trend.last_month_score:.0f}"
                month_background = "#164737" if month_text.startswith("↑") else "#4A2730" if month_text.startswith("↓") else PANEL_ALT
                ctk.CTkLabel(table, text=domain, anchor="w", font=("Segoe UI", 14)).grid(row=row, column=0, sticky="ew", padx=12, pady=6)
                ctk.CTkLabel(table, text=this_month_score, anchor="e", font=("Segoe UI", 14, "bold")).grid(row=row, column=1, sticky="ew", padx=12, pady=6)
                ctk.CTkLabel(table, text=last_month_score, anchor="e", font=("Segoe UI", 14)).grid(row=row, column=2, sticky="ew", padx=12, pady=6)
                ctk.CTkLabel(
                    table,
                    text=month_text,
                    width=82,
                    corner_radius=7,
                    fg_color=month_background,
                    anchor="e",
                    text_color=month_color,
                    font=("Segoe UI", 16, "bold"),
                ).grid(row=row, column=3, sticky="e", padx=12, pady=6)
                ctk.CTkLabel(table, text=year_text, anchor="e", text_color=year_color, font=("Segoe UI", 14, "bold")).grid(row=row, column=4, sticky="ew", padx=12, pady=6)

        switcher.configure(command=render)
        grade_menu.configure(command=render)
        # Variable traces pass (name, index, mode), which render() does not accept.
        search_text.trace_add("write", lambda *_: render())
        render()

    @staticmethod
    def _grade_for_score(score: float) -> str:
        if score >= 90:
            return "A"
        if score >= 80:
            return "B"
        if score >= 70:
            return "C"
        if score >= 60:
            return "D"
        return "F"

    def _history_by_month(self) -> dict[str, list[Company]]:
        """Read one latest score snapshot per domain for every workbook month."""
        latest: dict[tuple[int, int, str], tuple[datetime, Company]] = {}
        workbook = load_workbook(SCORE_HISTORY_PATH, read_only=True, data_only=True)
        try:
            for timestamp, domain, company_name, score, grade in workbook.active.iter_rows(
                min_row=3, max_col=5, values_only=True
            ):
                if not isinstance(timestamp, datetime) or not domain:
                    continue
                try:
                    numeric_score = float(score)
                except (TypeError, ValueError):
                    continue
                domain_text = str(domain).strip().lower()
                if not domain_text:
                    continue
                snapshot = timestamp.replace(tzinfo=None)
                key = (snapshot.year, snapshot.month, domain_text)
                record = Company(
                    domain=domain_text,
                    name=str(company_name or domain_text).strip(),
                    grade=str(grade or self._grade_for_score(numeric_score)).upper(),
                    score=numeric_score,
                )
                if key not in latest or snapshot > latest[key][0]:
                    latest[key] = (snapshot, record)
        finally:
            workbook.close()
        months: dict[str, list[Company]] = {}
        for (year, month, _domain), (_timestamp, company) in latest.items():
            label = datetime(year, month, 1).strftime("%B %Y")
            months.setdefault(label, []).append(company)
        return dict(sorted(months.items(), key=lambda item: datetime.strptime(item[0], "%B %Y"), reverse=True))

    def open_report_page(self) -> None:
        """Load score-history snapshots and open the monthly Report page."""
        if self.busy:
            return

        def build_report() -> None:
            snapshots = self._history_by_month()
            if not snapshots:
                raise ValueError("No monthly score snapshots are available in the score-history workbook.")
            self.after(0, lambda: self._show_report_page(snapshots))
            self._log(f"✓ Report loaded with {len(snapshots)} month(s) from {SCORE_HISTORY_PATH.name}.")

        self._run("Loading report history", build_report, show_loading_overlay=False)

    def _show_report_page(self, snapshots: dict[str, list[Company]]) -> None:
        """Render the workbook-backed monthly grade heatmap."""
        if self.trends_page is not None:
            self.trends_page.destroy()
            self.trends_page = None
        if self.report_page is not None:
            self.report_page.destroy()
        self.dashboard_header.grid_remove()
        self.dashboard_body.grid_remove()

        self.report_page = ctk.CTkScrollableFrame(self, fg_color=BACKGROUND, corner_radius=0)
        self.report_page.grid(row=0, column=1, rowspan=2, sticky="nsew", padx=(31, 23), pady=(0, 20))
        self.report_page.grid_columnconfigure(0, weight=1)
        ctk.CTkButton(
            self.report_page, text="←  Dashboard", command=self.show_dashboard,
            width=170, height=32, fg_color="transparent", hover_color=PANEL_ALT,
            anchor="w", font=("Segoe UI", 14),
        ).grid(row=0, column=0, sticky="w", padx=4, pady=(22, 0))
        heading = ctk.CTkFrame(self.report_page, fg_color="transparent")
        heading.grid(row=1, column=0, sticky="ew", padx=4, pady=(13, 11))
        heading.grid_columnconfigure(0, weight=1)
        ctk.CTkLabel(heading, text="Report", font=("Segoe UI", 25, "bold")).grid(row=0, column=0, sticky="w")
        ctk.CTkLabel(heading, text="View score distribution", font=("Segoe UI", 13), text_color=MUTED).grid(row=1, column=0, sticky="w", pady=(1, 0))
        period = ctk.StringVar(value=next(iter(snapshots)))
        selector = ctk.CTkOptionMenu(
            heading, values=list(snapshots), variable=period, width=180,
            font=("Segoe UI", 13), dropdown_font=("Segoe UI", 13),
        )
        selector.grid(row=0, column=1, rowspan=2, sticky="e")

        content = ctk.CTkFrame(self.report_page, fg_color=PANEL, corner_radius=14)
        content.grid(row=3, column=0, sticky="ew", padx=4, pady=(0, 8))
        content.grid_columnconfigure((0, 1, 2, 3, 4), weight=1, uniform="grades")
        for column, grade in enumerate(("A", "B", "C", "D", "F")):
            label, background, foreground = REPORT_GRADE_STYLES[grade]
            ctk.CTkLabel(
                content, text=label, height=52, fg_color=background, text_color=foreground,
                font=("Segoe UI", 14, "bold"), corner_radius=0,
            ).grid(row=0, column=column, sticky="ew", padx=(1 if column else 0, 0), pady=0)
        groups = [ctk.CTkFrame(content, fg_color="#0E1821", corner_radius=0) for _ in range(5)]
        for column, group in enumerate(groups):
            group.grid(row=1, column=column, sticky="nsew", padx=(1 if column else 0, 0), pady=(1, 0))
            group.grid_columnconfigure(0, weight=1)

        def render_month(_selection: str = "") -> None:
            grouped = {grade: [] for grade in REPORT_GRADE_STYLES}
            for company in sorted(snapshots[period.get()], key=lambda item: item.domain.lower()):
                grouped[company.grade if company.grade in grouped else self._grade_for_score(float(company.score or 0))].append(company)
            for index, grade in enumerate(("A", "B", "C", "D", "F")):
                group = groups[index]
                for widget in group.winfo_children():
                    widget.destroy()
                members = grouped[grade]
                if not members:
                    ctk.CTkLabel(group, text="Tidak ada domain", font=("Segoe UI", 12, "italic"), text_color=MUTED).grid(
                        row=0, column=0, sticky="w", padx=15, pady=17
                    )
                    continue
                for row, company in enumerate(members):
                    ctk.CTkLabel(
                        group, text=company.domain, anchor="w", height=24,
                        font=("Segoe UI", 13, "bold"), text_color="#5E9BFF",
                    ).grid(
                        row=row, column=0, sticky="ew", padx=15, pady=(7 if row == 0 else 2, 0)
                    )

        selector.configure(command=render_month)
        render_month()

    def show_dashboard(self) -> None:
        """Return from an in-window page to the main dashboard."""
        if self.trends_page is not None:
            self.trends_page.destroy()
            self.trends_page = None
        if self.report_page is not None:
            self.report_page.destroy()
            self.report_page = None
        self.dashboard_header.grid()
        self.dashboard_body.grid()

    def close_score_trends(self) -> None:
        """Compatibility alias for the score-trend back control."""
        self.show_dashboard()

    def start_score_export(self) -> None:
        if self.busy:
            return
        self._export_score_history()

    def _export_score_history(self) -> None:
        def export() -> None:
            client, portfolio_id = self._client()
            portfolio_entries = client.fetch_portfolio_companies(portfolio_id)
            selected_domains = {
                company.domain
                for company in companies_for_cycle_action(
                    companies_from_payload(portfolio_entries), self.active_cycle
                )
            }
            companies: list[Company] = []
            for item in portfolio_entries:
                domain = str(item.get("domain") or item.get("website") or "").strip()
                if not domain or normalize_domain(domain) not in selected_domains:
                    continue
                details = client.get_company(domain)
                companies.append(
                    Company(
                        str(details.get("domain") or domain),
                        str(details.get("name") or item.get("name") or domain),
                        str(details.get("grade") or "").upper(),
                        details.get("score"),
                    )
                )
            appended = append_score_history(companies, self.active_cycle or "Live portfolio")
            self.after(0, self._update_history_status)
            scope = "for the live portfolio" if self.active_cycle in (None, 1) else "excluding pinned domains"
            self._log(f"✓ Added {appended} score records {scope} to {SCORE_HISTORY_PATH}.")
        self._run("Updating live portfolio score history", export)


def main() -> None:
    app = Dashboard()
    logging.getLogger().handlers.clear()
    logging.getLogger().addHandler(QueueLogHandler(app.log_queue))
    logging.getLogger().setLevel(logging.INFO)
    app.mainloop()


if __name__ == "__main__":
    main()
